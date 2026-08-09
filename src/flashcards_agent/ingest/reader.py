"""Lectores de material fuente PDF/DOCX (TDD §4.4)."""

from __future__ import annotations

from pathlib import Path
from typing import Literal

import docx
import pymupdf

from flashcards_agent.models.content import PageText, SourceDocument


def read_pdf(path: Path) -> tuple[PageText, ...]:
    document = pymupdf.open(path)
    try:
        return tuple(
            PageText(
                number=i + 1,
                text=page.get_text(),
                image_count=len(page.get_images(full=True)),
            )
            for i, page in enumerate(document)
        )
    finally:
        document.close()


def read_docx(path: Path) -> tuple[str, ...]:
    document = docx.Document(path)
    return tuple(p.text.strip() for p in document.paragraphs if p.text.strip())


def _marker(filename: str) -> str:
    # Convención observada en el material real: "<Materia> - Semana N - EJ_x.ext" /
    # "... - R_x.ext" para ejercicios/clave; "<Materia> - Semana N.ext" (sin tercer
    # segmento) para teoría.
    stem = Path(filename).stem
    parts = stem.split(" - ")
    return parts[-1] if len(parts) >= 3 else ""


def _classify_kind(filename: str) -> Literal["theory", "exercises", "answer_key"]:
    marker = _marker(filename)
    if marker.startswith("EJ_"):
        return "exercises"
    if marker.startswith("R_"):
        return "answer_key"
    return "theory"


def exercise_key(filename: str) -> str:
    """Clave de pareo EJ_x <-> R_x (IPL-31) — una semana puede traer más de un par.

    "Mate - Semana 1 - EJ_1.1.docx" -> "1.1"; "Mate - Semana 1 - R_1.1.pdf" -> "1.1";
    "Nivelación Matemática - Semana 1.pdf" (teoría, sin marcador) -> "".
    """
    marker = _marker(filename)
    if marker.startswith("EJ_"):
        return marker.removeprefix("EJ_")
    if marker.startswith("R_"):
        return marker.removeprefix("R_")
    return ""


def discover_week(
    subject_slug: str, week: int, root: Path = Path("material")
) -> list[SourceDocument]:
    week_dirs = sorted(root.glob(f"*/{subject_slug}/semana-{week}"))
    if not week_dirs:
        return []

    documents: list[SourceDocument] = []
    for file_path in sorted(week_dirs[0].iterdir()):
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            pages = read_pdf(file_path)
        elif suffix == ".docx":
            paragraphs = read_docx(file_path)
            pages = tuple(
                PageText(number=i + 1, text=text, image_count=0)
                for i, text in enumerate(paragraphs)
            )
        else:
            continue

        documents.append(
            SourceDocument(
                path=file_path,
                kind=_classify_kind(file_path.name),
                subject_slug=subject_slug,
                week=week,
                pages=pages,
            )
        )
    return documents
